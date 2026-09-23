import io,zipfile
import pytest
from openpyxl import Workbook,load_workbook
from shipping.agent.ooxml import patch_xlsx,members


def test_sheet_selection_removes_other_sheets_and_unused_strings():
 from shipping.agent.document_ooxml import select_sheets
 w=Workbook();w.active.title='发票';w.active['A1']='CURRENT';w.create_sheet('报关')['A1']='PRIVATE_PRICE'
 b=io.BytesIO();w.save(b);out=select_sheets(b.getvalue(),['发票'])
 assert load_workbook(io.BytesIO(out)).sheetnames==['发票']
 assert b'PRIVATE_PRICE' not in b''.join(members(out).values())


def test_batch_sheet_does_not_change_existing_sheet_or_images():
 from shipping.agent.document_ooxml import append_batch_sheet
 from test_document_workflow import sample
 w=Workbook();w.active['A1']='FIXED';b=io.BytesIO();w.save(b);j=sample();out=append_batch_sheet(b.getvalue(),j,j['products'][0])
 before=members(b.getvalue());after=members(out)
 assert before['xl/worksheets/sheet1.xml']==after['xl/worksheets/sheet1.xml']
 wb=load_workbook(io.BytesIO(out));assert '批次明细' in wb.sheetnames
 assert any(c.value=='B1' for row in wb['批次明细'] for c in row)


def test_multi_batch_invoice_refers_to_batch_detail_not_first_date():
 from shipping.agent.profiles import customer_changes
 from test_document_workflow import sample
 j=sample();p=j['products'][0];p['batches'].append({'id':'b2','values':{'batch_no':'B2','mfg_date':'2026-08','exp_date':'2029-07'},'evidence':{},'conflicts':{}})
 changes=customer_changes(j,p)
 assert 'SEE BATCH DETAILS' in changes['发票!A23']
 assert 'B1' in changes['箱单!B23'] and 'B2' in changes['箱单!B23']
 assert '2026-08' not in changes['发票!A23']


def test_formula_quantity_uses_current_price_and_no_float():
 from shipping.agent.profiles import customer_changes
 from test_document_workflow import sample
 j=sample();p=j['products'][0];p['values'].update(quantity='3',customer_unit_price='0.1',customer_currency='EUR')
 assert customer_changes(j,p)['发票!L25']['value']=='0.30'


def test_attached_originals_require_explicit_selection_and_deduplicate_names(tmp_path):
 from shipping.agent.packages import attachment_package
 from shipping.repository import Repository
 repo=Repository(tmp_path);a=repo.put_file('证书.pdf',b'first');b=repo.put_file('证书.pdf',b'second');j={'files':[a,b],'attachment_selection':[]}
 with pytest.raises(ValueError):attachment_package(repo,j)
 j['attachment_selection']=[a['id'],b['id']];data=attachment_package(repo,j)
 with zipfile.ZipFile(io.BytesIO(data)) as z:
  assert len(z.namelist())==3 and len(set(z.namelist()))==3
  assert sorted(z.read(n) for n in z.namelist() if n.endswith('.pdf'))==[b'first',b'second']


def test_pickup_template_output_removes_signature_asset():
 from shipping.agent.document_ooxml import unsigned_pickup
 from shipping.agent.ooxml import W,R
 from lxml import etree as E
 from docx import Document
 from PIL import Image
 d=Document();img=io.BytesIO();Image.new('RGBA',(20,20),(200,0,0,255)).save(img,format='PNG');img.seek(0);d.add_picture(img);b=io.BytesIO();d.save(b)
 out=unsigned_pickup(b.getvalue());parts=members(out)
 assert not any(n.startswith('word/media/') for n in parts)
 assert not E.fromstring(parts['word/document.xml']).findall('.//{'+W+'}drawing')


def test_advice_does_not_block_on_irrelevant_sample_batch_capacity():
 from shipping.agent.profiles import customer_changes
 from test_document_workflow import sample
 j=sample();p=j['products'][0];p['values']['sample_batches']='A;B;C;D'
 changes=customer_changes(j,p,'advice')
 assert changes['SHIPPING ADVICE!B3']=='I1'
 with pytest.raises(ValueError):customer_changes(j,p,'packing')


def test_booking_print_fit_preserves_text_and_rejects_nonempty_extra_column():
 from shipping.agent.document_ooxml import fit_booking_tables
 from shipping.agent.ooxml import W
 from docx import Document
 from lxml import etree as E
 d=Document();outer=d.add_table(rows=2,cols=1)
 for row in outer.rows:
  t=row.cells[0].add_table(rows=2,cols=8)
  for c in t.rows[0].cells[:6]:c.text='固定表头'
 b=io.BytesIO();d.save(b);before=E.fromstring(members(b.getvalue())['word/document.xml']);out=fit_booking_tables(b.getvalue());after=E.fromstring(members(out)['word/document.xml'])
 assert ''.join(before.itertext())==''.join(after.itertext())
 for table in after.findall('.//{'+W+'}tbl')[1:]:
  assert len(table.findall('./{'+W+'}tblGrid/{'+W+'}gridCol'))==6
  assert all(len(row.findall('{'+W+'}tc'))==6 for row in table.findall('{'+W+'}tr'))
 t.rows[1].cells[7].text='不能删除';b=io.BytesIO();d.save(b)
 with pytest.raises(ValueError):fit_booking_tables(b.getvalue())
