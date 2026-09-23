from copy import deepcopy
from decimal import Decimal
import io,re
import pytest
from pypdf import PdfWriter
from shipping.agent.catalog import input_hash
from test_document_workflow import sample


def complete_job():
 j=sample();j['name']='QA';j['manual_entry']=True;j['manual_files']=[]
 j['contracts'][0]['values'].update(customs_contract_no='C1',customs_invoice_no='CI1',customs_invoice_date='2026-09-20',customs_buyer='QA BUYER')
 j['contracts'][0]['values']['payment_terms']='TT 60 DAYS FROM TRANSPORT DOCUMENT DATE'
 j['products'][0]['values'].update(customer_currency='USD',customer_unit_price='0.06',packing_text='25 TABS/BOTTLE',net_kg='10',cartons='4',storage='BELOW 25 C',samples_text='无',moc_no='无',insurance_no='无',customs_quantity='480',customs_unit='BOTTLE',declaration_unit_price='1.125',customs_currency='USD',customs_hs_code='3004909099',origin_country='CHINA',domestic_source='徐州',samples_declaration='无',declaration_elements='无',customs_unit_price='0.045',carton_net_kg='2.5',carton_gross_kg='3',chargeable_kg='15',transport_document_no='AWB-QA1',certificate_no='CO-QA1',certificate_date='2026-09-20')
 j['products'][0]['batches'][0]['values'].update(mfg_date='2026-01',exp_date='2029-01')
 j['products'][0]['values'].update(customs_tax_method='照章征税',customs_package_count='1',customs_main_net_kg='10')
 j['products'][0]['values']['certificate_hs_code']='3004'
 j['shipment']['values'].update(transport_mode='空运',loading_port='SHANGHAI',destination_port='TEHRAN',destination_country='IRAN',departure_port='SHANGHAI',trade_country='HONG KONG',trade_term='CPT',customs_trade_term='C&F',customs_freight='100',customs_insurance='0',customs_other_fee='0',fees_currency='USD',customs_packaging='托盘',customs_marks='N/M',customs_notes='无',departure_date='2026-09-21',flight='QA1',flight_date='2026-09-21',shipping_date='2026-09-20',warehouse_date='2026-09-21',freight_note='PREPAID')
 return j


def test_customs_samples_are_separate_decimal_lines_and_require_review():
 from shipping.agent.business_packages import customs_lines,save_samples,package_issues
 j=complete_job();j['products'][0]['values']['samples_declaration']='测试样品'
 assert any('样品' in x for x in package_issues(j,'customs_set'))
 j=save_samples(j,'p1',[dict(name='测试样品',quantity='2',unit='BOTTLE',unit_price='1',net_kg='0.00015',hs_code='3004909099',origin='CHINA',source='徐州',tax='照章征税',purpose='不收汇不退税')])
 lines=customs_lines(j)
 assert sum(x['amount'] for x in lines)==Decimal('542.00')
 assert sum(x['net'] for x in lines)==Decimal('10.00015')
 assert lines[1]['sample'] is True
 old=input_hash(j,'customs_set');j['customs_samples']['p1'][0]['quantity']='3'
 assert input_hash(j,'customs_set')!=old


def test_samples_cannot_disappear_after_toggle_or_accept_invalid_values():
 from shipping.agent.business_packages import save_samples,package_issues
 j=complete_job()
 with pytest.raises(ValueError):save_samples(j,'p1',[dict(name='S',quantity='NaN')])
 j['customs_samples']={'p1':[{'name':'S'}]}
 assert any('样品' in x for x in package_issues(j,'customs_set'))


def test_settlement_requires_each_products_originals_and_invalidates_selection(tmp_path):
 from shipping.repository import Repository
 from shipping.agent.business_packages import save_returns,package_issues,validate_returns
 repo=Repository(tmp_path);j=complete_job();writer=PdfWriter();writer.add_blank_page(200,200);b=io.BytesIO();writer.write(b)
 f=repo.put_file('return.pdf',b.getvalue());j['attachments']=[f]
 assert len([i for i in package_issues(j,'settlement') if '回件' in i])==2
 selection=[dict(product_id='p1',role=role,file_id=f['id'],pages=[1],confirmed=True,reference=ref) for role,ref in [('transport','AWB-QA1'),('certificate','CO-QA1')]]
 # One page cannot serve as both a transport document and a certificate.
 with pytest.raises(ValueError):save_returns(repo,j,selection)
 writer=PdfWriter();writer.add_blank_page(210,210);b=io.BytesIO();writer.write(b)
 f2=repo.put_file('co.pdf',b.getvalue());j['attachments'].append(f2);selection[1]['file_id']=f2['id']
 j=save_returns(repo,j,selection);assert not validate_returns(repo,j)
 h=input_hash(j,'settlement');j['returns'][0]['confirmed']=False
 assert input_hash(j,'settlement')!=h and package_issues(j,'settlement')


def test_group_multi_product_invoice_and_mixed_currency_block():
 from shipping.agent.business_packages import contract_groups,package_issues
 j=complete_job();p=deepcopy(j['products'][0]);p['id']='p2';j['products'].append(p)
 assert len(contract_groups(j))==1 and len(contract_groups(j)[0][1])==2
 p['values']['customer_currency']='EUR'
 assert any('币种' in x for x in package_issues(j,'shipping_draft'))


def test_fixed_forms_have_no_formulas_or_old_sample_values():
 from shipping.agent.business_forms import build_forms
 from openpyxl import load_workbook
 profile=dict(company_cn='测试公司',company_en='QA COMPANY',address='QA ADDRESS',consignee='QA CONSIGNEE',registration='QA-REG',seller_terms='QA TERMS')
 j=complete_job();j['products'][0]['values']['name_en']='=HYPERLINK("bad")'
 for kind,count in [('customs_set',4)]:
  parts=build_forms(j,kind,profile)
  assert len(parts)==count
  for part in parts:
   w=load_workbook(io.BytesIO(part['data']))
   assert len(w.sheetnames)==part['pages']
   assert not any(c.data_type=='f' for s in w for row in s for c in row)
   assert all(not s._images for s in w)


def test_shared_invoice_combines_contracts_once_without_losing_products():
 from shipping.agent.business_packages import contract_groups,package_issues
 j=complete_job();c=deepcopy(j['contracts'][0]);c['id']='c2';c['values']['contract_no']='C2';j['contracts'].append(c)
 p=deepcopy(j['products'][0]);p['id']='p2';p['contract_id']='c2';j['products'].append(p)
 groups=contract_groups(j);assert len(groups)==1 and groups[0][0]['values']['contract_no']=='C1 & C2'
 assert {p['id'] for p in groups[0][1]}=={'p1','p2'}
 c['values']['invoice_date']='2026-09-21';assert any('日期' in x for x in package_issues(j,'shipping_draft'))
 assert not any('日期' in x for x in package_issues(j,'consignment'))


def test_continuation_never_drops_long_paragraphs():
 from shipping.agent.business_forms import Form
 from shipping.agent.layout import normalized
 profile=dict(company_cn='测试公司',company_en='QA',address='TEST')
 text=' '.join('field'+str(i) for i in range(900))
 form=Form('续页',profile);form.paragraphs('DETAILS',[text])
 assert len(form.w.worksheets)>1
 assert normalized(text) in normalized(''.join(v for k,v in form.expected.items() if int(re.search(r'\d+$',k)[0])>=9))


def test_return_approval_expires_after_quantity_change(tmp_path):
 from shipping.repository import Repository
 from shipping.agent.business_packages import save_returns,package_issues
 repo=Repository(tmp_path);j=complete_job();rows=[]
 for i,(role,key) in enumerate([('transport','transport_document_no'),('certificate','certificate_no')]):
  w=PdfWriter();w.add_blank_page(200+i,200);b=io.BytesIO();w.write(b);f=repo.put_file(role+'.pdf',b.getvalue());j.setdefault('attachments',[]).append(f)
  rows.append(dict(product_id='p1',role=role,file_id=f['id'],pages=[1],confirmed=True,reference=j['products'][0]['values'][key]))
 j=save_returns(repo,j,rows);assert not package_issues(j,'settlement')
 j['products'][0]['values']['gross_kg']='16'
 assert len([x for x in package_issues(j,'settlement') if '回件' in x])==2


def test_page_range_order_and_out_of_range_are_explicit():
 from shipping.agent.business_packages import parse_pages
 assert parse_pages('2-3,1',4)==[2,3,1]
 assert parse_pages('',2)==[1,2]
 for bad in ['0','5','3-1','1,1','hello']:
  with pytest.raises(ValueError):parse_pages(bad,4)


def test_customs_packaging_and_net_weight_do_not_change_customer_weight():
 from shipping.agent.core import derive,resolve
 j=complete_job();p=j['products'][0];p['values'].pop('customs_main_net_kg');p['values'].pop('customs_package_count')
 derive(j);assert p['values']['customs_main_net_kg']=='10' and p['values']['customs_package_count']=='1'
 j=resolve(j,{'shipment.customs_packaging':'纸箱','p1.samples_declaration':'另列样品'},'本票包装核实')
 assert j['products'][0]['values']['customs_package_count']=='4'
 assert 'customs_main_net_kg' not in j['products'][0]['values']
 assert j['products'][0]['values']['net_kg']=='10'


def test_actual_package_renderer_all_four_and_original_hashes(tmp_path):
 import shutil,json,zipfile
 from pathlib import Path
 from shipping.repository import Repository
 from shipping.agent.business_packages import render_package,save_returns,package_current
 from shipping.ingestion import digest
 profiles=Path(__file__).parents[1]/'data'/'agent_profiles'
 if not (profiles/'business-v1.json').exists() or not (profiles/'co-v1.json').exists():pytest.skip('私有原件未安装；本机真实模板验收未运行')
 repo=Repository(tmp_path);shutil.copytree(profiles,repo.root/'agent_profiles')
 baseline={str(p.relative_to(repo.root)):digest(p.read_bytes()) for p in (repo.root/'agent_profiles').rglob('*') if p.is_file()}
 j=complete_job();rows=[]
 for i,(role,key) in enumerate([('transport','transport_document_no'),('certificate','certificate_no')]):
  w=PdfWriter();w.add_blank_page(200+i,200);b=io.BytesIO();w.write(b);f=repo.put_file(role+'.pdf',b.getvalue());j.setdefault('attachments',[]).append(f)
  rows.append(dict(product_id='p1',role=role,file_id=f['id'],pages=[1],confirmed=True,reference=j['products'][0]['values'][key]))
 j=save_returns(repo,j,rows)
 for kind in ['customs_set','shipping_draft','settlement','consignment']:
  out=render_package(repo,j,kind)[0];assert package_current(repo,out)
  with zipfile.ZipFile(repo.safe_path(out['path'])) as z:
   assert z.testzip() is None
   manifest=json.loads(z.read('文件清单.json'))
   for item in manifest['files']:assert digest(z.read(item['name']))==item['sha256']
   if kind in ('shipping_draft','settlement'):
    for item in manifest['files']:
     if item['name'].endswith('.xlsx'):
      with zipfile.ZipFile(io.BytesIO(z.read(item['name']))) as source:assert not any(n.startswith('xl/media/') for n in source.namelist())
   if kind=='settlement':
    assert len([x for x in manifest['components'] if x.get('original')])==2
    assert not any('草件' in x['name'] for x in manifest['components'])
  last=out
 assert baseline=={str(p.relative_to(repo.root)):digest(p.read_bytes()) for p in (repo.root/'agent_profiles').rglob('*') if p.is_file()}
 repo.safe_path(last['pdf_path']).write_bytes(b'changed');assert not package_current(repo,last)


def test_failed_package_does_not_leave_half_exports(tmp_path):
 import shutil
 from pathlib import Path
 from shipping.repository import Repository
 from shipping.agent.business_packages import render_package
 profiles=Path(__file__).parents[1]/'data'/'agent_profiles'
 if not (profiles/'business-v1.json').exists():pytest.skip('私有原件未安装')
 repo=Repository(tmp_path);shutil.copytree(profiles,repo.root/'agent_profiles');j=complete_job();j['products'][0]['values']['name_cn']='超长产品名称'*100
 with pytest.raises(ValueError,match='固定位置'):render_package(repo,j,'customs_set')
 assert not list((repo.root/'exports').rglob('*.xlsx')) and not list((repo.root/'exports').rglob('*.pdf'))


def test_public_file_pages_have_scoped_package_generation(tmp_path,monkeypatch):
 from streamlit.testing.v1 import AppTest
 from pathlib import Path
 from shipping.repository import Repository
 from shipping.agent.core import JobStore
 from shipping.agent import credentials
 monkeypatch.setenv('SHIPPING_DATA_DIR',str(tmp_path));monkeypatch.setattr(credentials,'available',lambda:False)
 store=JobStore(Repository(tmp_path));j=store.create('四套流程测试',['customs_set']);data=complete_job()
 for key in ('contracts','products','shipment','manual_entry','manual_files'):j[key]=data[key]
 j=store.save(j,j['revision'])
 app=AppTest.from_file(str(Path(__file__).parents[1]/'app.py'),default_timeout=30).run()
 app.radio(key='agent_step_'+j['id']).set_value('文件').run()
 for kind in ['customs_set','shipping_draft','settlement','consignment']:
  app.selectbox(key='agent_document_'+j['id']).select(kind).run();assert not app.exception
  button=app.button(key='package_generate_'+j['id']+'_'+kind)
  assert button.disabled==(kind=='settlement')
  assert not any('尚未接入' in x.value for x in app.caption)


def test_native_customer_reads_current_route_and_payment_from_their_cells(tmp_path):
 from openpyxl import Workbook
 from shipping.repository import Repository
 from shipping.agent.core import JobStore,normalise_result
 from shipping.agent.native import recognise
 from shipping.agent.flow import prepare_sources
 w=Workbook();s=w.active;s.title='发票';w.create_sheet('箱单');w.create_sheet('SHIPPING ADVICE')
 for k,v in {'A4':'INVOICE','L7':'INV-QA','L10':'2026-09-20','D23':'TEST TABLETS','G36':'TT 90 DAYS FROM RECEIPT','C14':'HORGOS','I14':'TEHRAN','K21':'CPT BY TRUCK'}.items():s[k]=v
 w['SHIPPING ADVICE']['B27']='TT 90 DAYS FROM RECEIPT';b=io.BytesIO();w.save(b)
 repo=Repository(tmp_path);f=repo.put_file('customer.xlsx',b.getvalue());job=JobStore(repo).create('QA',['shipping_draft']);job['files']=[f];sources=prepare_sources(repo,job)
 raw=recognise(repo,job,sources)
 assert any(x['key']=='payment_terms' and x['value']=='TT 90 DAYS FROM RECEIPT' for x in raw['contracts'][0]['fields'])
 assert any(x['key']=='transport_mode' and x['value']=='陆运' for x in raw['shipment'])


def test_pdfium_line_end_hyphen_keeps_its_character_and_visible_ink():
 from pypdf.generic import NameObject,DictionaryObject,DecodedStreamObject
 from shipping.agent.layout import check_pdf
 writer=PdfWriter();page=writer.add_blank_page(300,300)
 font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
 page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
 stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 30 250 Td (QA-) Tj 0 -15 Td (INS-001) Tj ET');page[NameObject('/Contents')]=writer._add_object(stream)
 b=io.BytesIO();writer.write(b)
 assert check_pdf(b.getvalue(),{'insurance':'QA-INS-001'},1)['status']=='passed'
