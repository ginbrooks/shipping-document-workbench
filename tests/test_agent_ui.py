from pathlib import Path
from streamlit.testing.v1 import AppTest
import pytest


def test_unexpected_extraction_failure_stops_card_and_keeps_saved_job(tmp_path,monkeypatch):
 from shipping.services import Workbench
 from shipping.agent.core import JobStore
 from shipping.ui import agent
 root=tmp_path/'data';monkeypatch.setenv('SHIPPING_DATA_DIR',str(root))
 w=Workbench(root);store=JobStore(w.repo);job=store.create('异常恢复',['booking'])
 job['files']=[w.repo.put_file('test.png',b'synthetic')];job=store.save(job,job['revision'])
 monkeypatch.setattr(agent,'read_with_progress',lambda *a,**kw:{'f:1':'synthetic'})
 def fail(*a,**kw):raise TypeError('SENSITIVE_CONTENT_MUST_NOT_BE_LOGGED')
 monkeypatch.setattr(agent,'extract_and_clean',fail)
 monkeypatch.setattr(agent,'connection_key',lambda endpoint:'synthetic-key')
 app=AppTest.from_file(str(Path(__file__).parents[1]/'app.py'),default_timeout=20).run()
 app.radio(key='agent_step_'+job['id']).set_value('资料').run()
 app.button(key='agent_extract_'+job['id']).click().run()
 assert not app.exception
 assert any('程序错误' in e.value and '错误编号' in e.value for e in app.error)
 assert any('charge-panel error' in e.value and 'aria-busy="false"' in e.value for e in app.markdown)
 assert store.get(job['id'])==job
 diagnostic=w.repo.cache_get('agent-last-processing-error')
 assert diagnostic['error_type']=='TypeError' and diagnostic['frames']
 import json
 assert 'SENSITIVE_CONTENT' not in json.dumps(diagnostic)

@pytest.fixture(autouse=True)
def isolate_system_credentials(monkeypatch):
 from shipping.agent import credentials
 monkeypatch.setattr(credentials,"available",lambda:False)
 monkeypatch.setattr(credentials,"load_key",lambda endpoint:None)


def test_agent_can_start_without_legacy_shipment(tmp_path,monkeypatch):
 monkeypatch.setenv('SHIPPING_DATA_DIR',str(tmp_path/'data'))
 app=AppTest.from_file(str(Path(__file__).parents[1]/'app.py'),default_timeout=20).run()
 assert not app.exception
 app.text_input(key='review_name_new').set_value('测试自动出单')
 app.button(key='review_save_new').click().run()
 assert not app.exception and not app.error
 assert app.selectbox(key='agent_selected').value

def test_intake_scope_updates_for_added_files_without_automatic_send(tmp_path,monkeypatch):
 import io
 from docx import Document
 from shipping.services import Workbench
 from shipping.agent.core import JobStore
 from shipping.agent.flow import prepare_sources
 from shipping.ui import agent
 root=tmp_path/'data';monkeypatch.setenv('SHIPPING_DATA_DIR',str(root))
 w=Workbench(root);store=JobStore(w.repo);j=store.create('读取进度测试',['customer'])
 doc=Document();doc.add_paragraph('测试合同资料');buf=io.BytesIO();doc.save(buf)
 j['files']=[w.repo.put_file('test.docx',buf.getvalue())];j=store.save(j,j['revision']);prepare_sources(w.repo,j)
 calls=[];monkeypatch.setattr(agent,'extract_and_clean',lambda *a,**kw:calls.append(True))
 app=AppTest.from_file(str(Path(__file__).parents[1]/'app.py'),default_timeout=30).run()
 assert not app.exception and not app.error and not calls
 assert any('test.docx' in e.value for e in app.markdown)
 assert not any(str(c.key).startswith('agent_send_') for c in app.checkbox)
 assert w.repo.cache_get('agent-last-api-diagnostic') is None
 doc.add_paragraph('新增资料');buf=io.BytesIO();doc.save(buf)
 updated=store.get(j['id']);updated['files'].append(w.repo.put_file('extra.docx',buf.getvalue()));store.save(updated,updated['revision'])
 app.run()
 assert any('extra.docx' in e.value for e in app.markdown)
 assert not calls

def test_keychain_ui_reloads_in_new_session_and_clears_on_endpoint_change(tmp_path,monkeypatch):
 from shipping.agent import credentials
 from shipping.services import Workbench
 from shipping.agent.core import JobStore
 saved={}
 monkeypatch.setattr(credentials,'available',lambda:True)
 monkeypatch.setattr(credentials,'load_key',lambda endpoint:saved.get(endpoint))
 monkeypatch.setattr(credentials,'save_key',lambda endpoint,key:saved.update({endpoint:key}))
 monkeypatch.setattr(credentials,'delete_key',lambda endpoint:saved.pop(endpoint,None))
 root=tmp_path/'data';monkeypatch.setenv('SHIPPING_DATA_DIR',str(root))
 w=Workbench(root);JobStore(w.repo).create('密钥保存测试',['customer'])
 w.repo.cache_set('agent-public-config',{'endpoint':'https://api.deepseek.com','model':'test-model'})
 path=str(Path(__file__).parents[1]/'app.py')
 app=AppTest.from_file(path,default_timeout=20).run()
 app.radio(key='nav_page').set_value('设置').run()
 assert any('输入 API 密钥' in x.value for x in app.info)
 app.text_input(key='agent_api_key').set_value('synthetic-ui-key').run()
 app.button(key='agent_remember_key').click().run()
 assert not app.exception and saved['https://api.deepseek.com']=='synthetic-ui-key'
 fresh=AppTest.from_file(path,default_timeout=20).run()
 fresh.radio(key='nav_page').set_value('设置').run()
 assert fresh.text_input(key='agent_api_key').value=='synthetic-ui-key'
 fresh.text_input(key='agent_endpoint').set_value('https://different.invalid').run()
 assert fresh.text_input(key='agent_api_key').value==''
 fresh.text_input(key='agent_endpoint').set_value('https://api.deepseek.com').run()
 assert fresh.text_input(key='agent_api_key').value=='synthetic-ui-key'
 fresh.button(key='agent_forget_key').click().run()
 assert not fresh.exception and not saved and fresh.text_input(key='agent_api_key').value==''
 for file in root.rglob('*'):
  if file.is_file():assert b'synthetic-ui-key' not in file.read_bytes()

def test_cleaned_fields_are_prefilled_and_saving_one_does_not_mark_all_manual(tmp_path,monkeypatch):
 from shipping.services import Workbench
 from shipping.agent.core import JobStore
 root=tmp_path/'data';monkeypatch.setenv('SHIPPING_DATA_DIR',str(root))
 w=Workbench(root);store=JobStore(w.repo);j=store.create('自动回填验证',['customer'])
 def obj(oid,values,**extra):return {'id':oid,'values':values,'evidence':{k:[{'source':'source:1','quote':v,'value':v}] for k,v in values.items()},'conflicts':{},**extra}
 j['contracts']=[obj('c1',{'contract_no':'C001'},kind='purchase')]
 j['products']=[obj('p1',{'name_cn':'测试片','name_en':'TEST','carton_length_mm':'390'},contract_id='c1',batches=[obj('b1',{'batch_no':'B001','mfg_date':'2026-08-07','exp_date':'2029-07'})])]
 j['phase']='needs_input';j=store.save(j,j['revision'])
 app=AppTest.from_file(str(Path(__file__).parents[1]/'app.py'),default_timeout=20).run()
 app.radio(key='agent_step_'+j['id']).set_value('核对').run()
 prefix='af_'+j['id']+'_'+str(j['revision'])+'_'
 assert app.text_input(key=prefix+'p1.name_en').value=='TEST'
 app.radio(key='review_group_'+j['id']+'_shipping_draft').set_value('packing').run()
 assert app.text_input(key=prefix+'b1.mfg_date').value=='2026-08-07'
 assert app.text_input(key=prefix+'b1.exp_date').value=='2029-07'
 app.radio(key='review_group_'+j['id']+'_shipping_draft').set_value('shared').run()
 app.text_input(key=prefix+'c1.invoice_no').set_value('INV001')
 app.button(key='review_save_'+j['id']).click().run()
 assert not app.exception
 saved=store.get(j['id'])
 assert saved['contracts'][0]['values']['invoice_no']=='INV001'
 assert saved['products'][0]['evidence']['name_en'][0]['source']=='source:1'


def test_invalid_saved_key_is_not_reported_as_missing_and_failed_save_preserves_it(tmp_path,monkeypatch):
 from shipping.agent import credentials
 from shipping.services import Workbench
 from shipping.agent.core import JobStore
 saved={'https://api.deepseek.com':'无效的旧内容'}
 monkeypatch.setattr(credentials,'available',lambda:True)
 monkeypatch.setattr(credentials,'load_key',lambda endpoint:saved.get(endpoint))
 monkeypatch.setattr(credentials,'save_key',lambda endpoint,key:saved.update({endpoint:key}))
 root=tmp_path/'data';monkeypatch.setenv('SHIPPING_DATA_DIR',str(root))
 w=Workbench(root);store=JobStore(w.repo);j=store.create('旧密钥恢复测试',['customer'])
 j['files']=[w.repo.put_file('test.png',b'test')];store.save(j,j['revision'])
 w.repo.cache_set('agent-public-config',{'endpoint':'https://api.deepseek.com','model':'test-model'})
 app=AppTest.from_file(str(Path(__file__).parents[1]/'app.py'),default_timeout=20).run()
 assert not app.exception
 app.radio(key='agent_step_'+j['id']).set_value('资料').run()
 assert any('修正已保存的无效密钥' in x.value for x in app.info)
 app.radio(key='nav_page').set_value('设置').run()
 assert any('已保存并读取，但格式无效' in x.value for x in app.warning)
 assert not any('已保存在本机钥匙串' in x.value for x in app.success)
 assert not any('输入 API 密钥' in x.value for x in app.info)
 app.button(key='agent_remember_key').click().run()
 assert saved['https://api.deepseek.com']=='无效的旧内容'
 app.text_input(key='agent_api_key').set_value('synthetic-repaired-key').run()
 assert any('尚未保存' in x.value for x in app.caption)
 app.button(key='agent_remember_key').click().run()
 assert not app.exception
 fresh=AppTest.from_file(str(Path(__file__).parents[1]/'app.py'),default_timeout=20).run()
 fresh.radio(key='nav_page').set_value('设置').run()
 assert fresh.text_input(key='agent_api_key').value=='synthetic-repaired-key'
 assert any('已保存在本机钥匙串' in x.value for x in fresh.success)
 assert fresh.session_state['agent_key_status']=='valid'
