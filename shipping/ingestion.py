"""Bounded, content-addressed import. Document contents never execute."""
import hashlib
import io
import stat
import zipfile
from pathlib import PurePosixPath

SUPPORTED = {'.xlsx', '.docx', '.pdf', '.png', '.jpg', '.jpeg', '.rar'}


def digest(data):
    return hashlib.sha256(data).hexdigest()


def archive_name(info):
    name = info.filename
    if not info.flag_bits & 0x800:
        raw = name.encode('cp437')
        for encoding in ('utf-8', 'gb18030'):
            try:
                return raw.decode(encoding)
            except UnicodeError:
                pass
    return name


def unpack(name, data, max_upload=100*1024**2, max_expanded=500*1024**2,
           max_entries=1000, max_depth=2):
    if len(data) > max_upload:
        raise ValueError('UPLOAD_TOO_LARGE')
    result, budget = [], {'bytes': 0, 'entries': 0}

    def walk(display, payload, depth):
        if PurePosixPath(display).suffix.lower() != '.zip':
            result.append({'name': display, 'data': payload, 'sha256': digest(payload),
                           'status': 'UNPARSED_RAR' if display.lower().endswith('.rar') else 'IMPORTED'})
            return
        if depth > max_depth:
            raise ValueError('ZIP_DEPTH_LIMIT')
        with zipfile.ZipFile(io.BytesIO(payload)) as z:
            for entry in z.infolist():
                n = archive_name(entry).replace('\\', '/')
                path = PurePosixPath(n)
                if path.is_absolute() or '..' in path.parts or ':' in path.parts[0] or stat.S_ISLNK(entry.external_attr >> 16):
                    raise ValueError('UNSAFE_ZIP_PATH')
                if entry.is_dir() or any(p == '__MACOSX' or p == '.DS_Store' or p.startswith('._') for p in path.parts):
                    continue
                budget['entries'] += 1
                budget['bytes'] += entry.file_size
                if budget['entries'] > max_entries or budget['bytes'] > max_expanded:
                    raise ValueError('ZIP_EXPANSION_LIMIT')
                if entry.flag_bits & 1:
                    raise ValueError('ENCRYPTED_ZIP_UNSUPPORTED')
                with z.open(entry) as f:
                    content = f.read(min(entry.file_size + 1, max_expanded + 1))
                if len(content) != entry.file_size:
                    raise ValueError('ZIP_SIZE_MISMATCH')
                walk(display + '/' + n, content, depth + 1)
    walk(name, data, 0)
    return result


def classify(name):
    n = name.lower()
    for words, role in ((('对照', 'reference'), 'reference_coa'), (('coa', '检验'), 'finished_coa'),
                        (('合同', 'contract'), 'contract'), (('提单', 'awb'), 'carrier_return'),
                        (('报关',), 'customs'), (('发票', 'invoice'), 'invoice')):
        if any(w in n for w in words):
            return role
    return 'unclassified'


def extract_local(name, data):
    """Explicit invocation only; callers persist results keyed by hash."""
    suffix = PurePosixPath(name).suffix.lower()
    entries = []
    if suffix == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        count = 0
        for ws in wb:
            if ws.max_row > 10000 or ws.max_column > 500:
                raise ValueError('SHEET_TOO_LARGE_SELECT_RELEVANT_RANGE')
            for row in ws:
                for c in row:
                    if c.value is not None:
                        entries.append({'locator': f'{ws.title}!{c.coordinate}', 'text': str(c.value)})
                        count += 1
                        if count > 30000:
                            raise ValueError('TEXT_LIMIT_SELECT_RELEVANT_RANGE')
        wb.close()
    elif suffix == '.docx':
        from docx import Document
        doc = Document(io.BytesIO(data))
        entries += [{'locator': f'paragraph:{i+1}', 'text': p.text} for i, p in enumerate(doc.paragraphs) if p.text.strip()]
        def read_tables(tables,prefix=''):
          for t, table in enumerate(tables):
            seen = set()
            for r, row in enumerate(table.rows):
                for c, cell in enumerate(row.cells):
                    if cell._tc not in seen:
                        seen.add(cell._tc);loc=prefix+f'table:{t+1}/row:{r+1}/cell:{c+1}'
                        if cell.text.strip():entries.append({'locator':loc,'text':cell.text})
                        if cell.tables:read_tables(cell.tables,loc+'/')
        read_tables(doc.tables)
    elif suffix == '.pdf':
        from pypdf import PdfReader
        for i, page in enumerate(PdfReader(io.BytesIO(data)).pages):
            value = page.extract_text() or ''
            entries.append({'locator': f'page:{i+1}', 'text': value,
                            'status': 'TEXT' if value.strip() else 'NOT_CHECKED_SCAN'})
    else:
        entries.append({'locator': 'file', 'text': '', 'status': 'UNPARSED_RAR' if suffix == '.rar' else 'NOT_CHECKED_IMAGE'})
    return {'entries': entries, 'role_candidate': classify(name), 'sha256': digest(data)}
