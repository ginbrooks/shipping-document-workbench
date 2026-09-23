"""Lossless readers for the fixed, labelled native templates. No guesses or network.
Use only for one recognised file; mixed evidence still uses the source-checked model.
"""
import io,re
from pathlib import Path
from openpyxl import load_workbook
from lxml import etree
from .ooxml import members,W


def recognise(repo,job,sources):
 if len(job['files'])!=1:return None
 file=job['files'][0];name=file['name'];prefix=file['id']+':'
 raw={'contracts':[{'id':'c1','kind':'other','fields':[]}],'products':[],'shipment':[]}
 def fact(fields,key,value,quote,locator=None):
  if value in (None,'','待补'):return
  source=prefix+locator if locator else next((s for s,t in sources.items() if s.startswith(prefix) and re.sub(r'\s+','',str(quote))==re.sub(r'\s+','',t)),None)
  if source not in sources:return
  fields.append({'key':key,'value':str(value).strip(),'source':source,'quote':str(quote).strip()})
 def number(fields,key,text,locator=None):
  m=re.search(r'(?<![\w.])\d+(?:,\d{3})*(?:\.\d+)?',str(text or ''))
  if m:fact(fields,key,m[0].replace(',',''),text,locator)
 if Path(name).suffix.lower()=='.xlsx':
  wb=load_workbook(io.BytesIO(repo.file_bytes(file['id'])),data_only=False)
  if not {'发票','箱单','SHIPPING ADVICE'}<=set(wb.sheetnames) or wb['发票']['A4'].value!='INVOICE':return None
  p={'id':'p1','contract_id':'c1','fields':[],'batches':[]};raw['products'].append(p)
  def read(key,loc,fields=None,numeric=False):
   sheet,cell=loc.split('!');text=wb[sheet][cell].value
   if text is not None and not str(text).startswith('='):(number if numeric else fact)(fields if fields is not None else p['fields'],key,text,loc) if numeric else fact(fields if fields is not None else p['fields'],key,text,text,loc)
  cf=raw['contracts'][0]['fields']
  read('invoice_no','发票!L7',cf);read('invoice_date','发票!L10',cf)
  read('payment_terms','发票!G36',cf);read('payment_terms','SHIPPING ADVICE!B27',cf)
  read('loading_port','发票!C14',raw['shipment']);read('destination_port','发票!I14',raw['shipment'])
  terms=str(wb['发票']['K21'].value or '')
  m=re.fullmatch(r'\s*([A-Z]{3})\s+BY\s+(AIR|TRUCK)\s*',terms,re.I)
  if m:
   fact(raw['shipment'],'trade_term',m[1].upper(),terms,'发票!K21');fact(raw['shipment'],'transport_mode','空运' if m[2].upper()=='AIR' else '陆运',terms,'发票!K21')
  # Invoice number does not automatically establish a sales contract number.
  read('name_en','发票!D23')
  packing=str(wb['发票']['F27'].value or '');fact(p['fields'],'packing_text',re.sub(r'(?i)^PACKAGE:\s*','',packing),packing,'发票!F27')
  read('quantity','发票!H23',numeric=True)
  unit=str(wb['发票']['H23'].value or '');m=re.search(r'(?i)\b(TABS?|BOTTLES?|KG|PCS)\b',unit)
  if m:fact(p['fields'],'base_unit',m[0],unit,'发票!H23')
  price=str(wb['发票']['I23'].value or '');m=re.search(r'\b([A-Z]{3})\s*([\d.]+)\s*/',price)
  if m:fact(p['fields'],'customer_currency',m[1],price,'发票!I23');fact(p['fields'],'customer_unit_price',m[2],price,'发票!I23')
  for k,loc in [('cartons','箱单!F15'),('net_kg','箱单!I15'),('gross_kg','箱单!H15'),('volume_m3','箱单!J15')]:read(k,loc,numeric=True)
  text=str(wb['箱单']['B23'].value or '');m=re.fullmatch(r'(?i)BATCH\s*NO[.:：\s]*([A-Za-z0-9_-]+)',text.strip())
  if m:fact(p['fields'],'batch_no',m[1],text,'箱单!B23')
  return raw
 if Path(name).suffix.lower()!='.docx':return None
 root=etree.fromstring(members(repo.file_bytes(file['id']))['word/document.xml']);tables=root.findall('.//{'+W+'}tbl')
 if not tables:return None
 def cell(t,r,c):
  try:return '\n'.join(''.join(n.text or '' for n in p.findall('.//{'+W+'}t')) for p in tables[t].findall('{'+W+'}tr')[r].findall('{'+W+'}tc')[c].findall('{'+W+'}p')).strip()
  except IndexError:return ''
 if len(tables)==1 and '取报告委托书' in ''.join(root.itertext()) and '报告编号' in cell(0,0,0):
  for i in range(1,len(tables[0].findall('{'+W+'}tr'))):
   report=cell(0,i,0);title=cell(0,i,1)
   if not title:continue
   p={'id':'p'+str(i),'contract_id':'c1','fields':[]};raw['products'].append(p)
   fact(p['fields'],'report_no',re.sub(r'(?i)^NO[. :]*','',report),report);fact(p['fields'],'name_cn',title,title)
   copies=cell(0,i,2);m=re.fullmatch(r'([0-9]+|壹|贰|叁|一|二|三)\s*份',copies)
   if m:fact(p['fields'],'report_copies',{'壹':'1','贰':'2','叁':'3','一':'1','二':'2','三':'3'}.get(m[1],m[1]),copies)
  for source,text in sources.items():
   if source.startswith(prefix) and re.fullmatch(r'\d{4}年\s*\d{1,2}月\s*\d{1,2}\s*日',text.strip()):fact(raw['shipment'],'authorization_date',text,text)
  return raw if raw['products'] else None
 if len(tables)!=3:return None
 if '出口货物托运单' not in re.sub(r'\s+','',cell(0,0,0)) or '外销实际单价' not in cell(1,0,2):return None
 text=cell(0,2,5);fact(raw['contracts'][0]['fields'],'contract_no',text,text)
 header=cell(0,0,0);m=re.search(r'\d{4}年\s*\d{1,2}月\s*\d{1,2}\s*日',header)
 if m:fact(raw['shipment'],'shipping_date',m[0],header)
 for key,pos in [('warehouse_date',(0,12,1)),('freight_note',(0,16,3))]:
  txt=cell(*pos);fact(raw['shipment'],key,txt,txt)
 for index in range(1,len(tables[1].findall('{'+W+'}tr'))):
  text=cell(1,index,0)
  if not text:continue
  p={'id':'p'+str(index),'contract_id':'c1','fields':[],'batches':[]};raw['products'].append(p)
  for key,label in [('name_cn','中文'),('name_en','英文')]:
   m=re.search(label+r'[：:]\s*([^\n]+)',text)
   if m:fact(p['fields'],key,m[1],text)
  text=cell(1,index,1);m=re.search(r'([\d,]+)\s*(TABS?|片|BOTTLES?|瓶|PCS)',text,re.I)
  if m:fact(p['fields'],'quantity',m[1].replace(',',''),text);fact(p['fields'],'base_unit',m[2],text)
  for key,pat in [('cartons',r'([\d,]+)\s*箱'),('pallets',r'([\d,]+)\s*托')]:
   m=re.search(pat,text)
   if m:fact(p['fields'],key,m[1].replace(',',''),text)
  for group,col in [('customer',2),('customs',4)]:
   txt=cell(1,index,col);m=re.search(r'\b([A-Z]{3})\s*([\d.]+)\s*/?',txt)
   if m:fact(p['fields'],group+'_currency',m[1],txt);fact(p['fields'],group+'_unit_price',m[2],txt)
  for key,col in [('packing_text',1),('net_kg',3),('carton_net_kg',4),('carton_gross_kg',5)]:
   txt=cell(2,index,col)
   if key=='packing_text':fact(p['fields'],key,txt,txt)
   else:number(p['fields'],key,txt)
  txt=cell(0,17,3)
  if re.fullmatch(r'\d{6,12}',txt):fact(p['fields'],'hs_code',txt,txt)
 return raw if raw['products'] else None
