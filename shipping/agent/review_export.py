"""Export the reviewed facts, never label a checklist as a filed customs document."""
from io import BytesIO
import unicodedata

def lines(text,capacity):
 result=[];line="";used=0
 for char in str(text):
  size=2 if unicodedata.east_asian_width(char) in ("W","F") else 1
  if char=="\n":result.append(line);line="";used=0;continue
  if used+size>capacity:result.append(line);line="";used=0
  line+=char;used+=size
 result.append(line);return result

from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from .business_flows import review_rows,COMPONENTS
from .catalog import NAMES

def workbook(job,kind):
 w=Workbook();s=w.active;s.title='核对信息';s.append([NAMES[kind]+' · '+job['name']]);s.append(['包含：'+'、'.join(COMPONENTS[kind])]);s.append(['本表用于核对数据，不代替报关、运输或证书正式文件。'])
 rows=review_rows(job,kind);headers=['归属','字段','核对值','原文依据','状态'];s.append(headers)
 for row in rows:
  wrapped=[lines(str(row.get(k) or ''),cap) for k,cap in zip(headers,(14,18,26,40,6))]
  for start in range(0,max(map(len,wrapped)),5):
   values=['\n'.join(v[start:start+5]) for v in wrapped]
   s.append(values);s.row_dimensions[s.max_row].height=max(1,max(min(5,len(v)-start) for v in wrapped))*14+6
 # Keep uploaded text as text, including formula-like strings.
 for row in s:
  for cell in row:
   cell.data_type='s';cell.font=Font(name='Arial',size=10);cell.alignment=Alignment(vertical='top',wrap_text=True,shrink_to_fit=False)
 for row in (s[1],s[4]):
  for c in row:c.font=Font(name='Arial',size=10,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='35536B')
 for k,v in {'A':18,'B':22,'C':32,'D':48,'E':8}.items():s.column_dimensions[k].width=v
 for n,height in ((1,30),(2,30),(3,30)):
  s.merge_cells(start_row=n,start_column=1,end_row=n,end_column=5);s.row_dimensions[n].height=height
 s.freeze_panes='C5';s.auto_filter.ref=f'A4:E{max(4,s.max_row)}';s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A4;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0;s.print_title_rows='1:4'
 b=BytesIO();w.save(b);return b.getvalue()
